"""Build the Word manual from README.md. Requires python-docx.

Only headings, paragraphs, code blocks, flat lists and pipe tables are used.
The public README is the editable text source; render the DOCX after changes.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def runs(paragraph, text, table=False):
    # Keep link labels readable; URLs remain available in the Markdown guide.
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    for part in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
        run = paragraph.add_run(part.strip("`") if part.startswith("`")
                                else part.strip("*") if part.startswith("**") else part)
        if part.startswith("`"):
            run.font.name = "Consolas"
            run.font.size = Pt(8.5 if table else 9)
        if part.startswith("**"):
            run.bold = True


def build():
    doc = Document()
    # Some bundled default templates put a colored border under Title.
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin = sec.bottom_margin = Inches(.65)
    sec.left_margin = sec.right_margin = Inches(.7)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.06
    doc.styles["Title"].font.size = Pt(25)
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(14)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(7)
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("NflowModel and VISIToRN   |   ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(8)
    lines = (ROOT / "README.md").read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith("# "):
            doc.add_paragraph(line[2:], "Title")
        elif line.startswith("## "):
            doc.add_paragraph(line[3:], "Heading 1")
        elif line.startswith("```"):
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1
                p.paragraph_format.keep_with_next = True
                run = p.add_run(lines[index])
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                index += 1
            p.paragraph_format.keep_with_next = False
            p.paragraph_format.space_after = Pt(7)
        elif line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].startswith("|"):
                row = [cell.strip() for cell in lines[index].strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+", cell) for cell in row):
                    rows.append(row)
                index += 1
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = "Table Grid"
            table.autofit = False
            widths = [2.1, 4.77] if len(rows[0]) == 2 else [2.2, 1.2, 3.47]
            if rows[0][0] == "Purpose":
                widths = [1.8, 2.5, 2.57]
            for col, width in zip(table.columns, widths):
                col.width = Inches(width)
            for row_index, row in enumerate(rows):
                cells = table.add_row().cells
                trpr = table.rows[-1]._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:cantSplit"))
                if row_index == 0:
                    trpr.append(OxmlElement("w:tblHeader"))
                for cell, text, width in zip(cells, row, widths):
                    cell.width = Inches(width)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.space_before = Pt(3)
                    runs(p, text, table=True)
                    for run in p.runs:
                        if run.font.size is None:
                            run.font.size = Pt(9.5)
                        if row_index == 0:
                            run.bold = True
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue
        else:
            bullet = line.startswith("- ")
            text = line[2:] if bullet else line
            while (index + 1 < len(lines) and lines[index + 1].strip()
                   and not lines[index + 1].startswith(("#", "|", "```", "- "))):
                index += 1
                text += " " + lines[index].strip()
            p = doc.add_paragraph(style="List Bullet" if bullet else "Normal")
            runs(p, text)
        index += 1
    doc.core_properties.title = "NflowModel and VISIToRN user guide"
    doc.core_properties.author = "Yuto Osugi"
    doc.core_properties.subject = "GMD reproduction and additional multi-scenario analysis"
    doc.save(ROOT / "README.docx")


if __name__ == "__main__":
    build()
